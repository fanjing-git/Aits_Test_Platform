"""Bounded, dependency-light parsers for requirement documents."""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from docx import Document as WordDocument
from pypdf import PdfReader

from apps.requirement_analysis.adapters import DocumentContent, SwaggerAdapter, adapter_for


MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".markdown", ".pdf", ".docx", ".xlsx", ".json", ".png", ".jpg", ".jpeg", ".webp"})


class DocumentParseError(ValueError):
    """Raised when a requirement source cannot be safely parsed."""


@dataclass(frozen=True)
class ParsedDocument:
    """Normalized parser output shared by uploaded and online sources."""

    title: str
    content: str
    format: str
    source_type: str
    source_url: str = ""
    evidence: tuple[dict[str, Any], ...] = ()
    confidence: float = 1.0
    warnings: tuple[str, ...] = ()


def _normalize(text: str) -> str:
    """Normalize line endings and remove empty trailing whitespace."""
    return "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")).strip()


def _bounded(content: bytes) -> bytes:
    if not content:
        raise DocumentParseError("文档内容不能为空。")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DocumentParseError("文档超过10MB限制。")
    return content


def _parse_xlsx(content: bytes) -> tuple[str, list[dict[str, Any]]]:
    """Extract visible cell values from XLSX XML without evaluating formulas."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            shared: list[str] = []
            if "xl/sharedStrings.xml" in names:
                root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
                shared = ["".join(node.text or "" for node in item.iter() if node.tag.endswith("}t") or node.tag == "t") for item in root if item.tag.endswith("}si") or item.tag == "si"]
            sheets = sorted(name for name in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name))
            rows: list[str] = []
            evidence: list[dict[str, Any]] = []
            for sheet in sheets:
                root = ElementTree.fromstring(archive.read(sheet))
                for row in root.iter():
                    if not row.tag.endswith("}row") and row.tag != "row":
                        continue
                    values: list[str] = []
                    for cell in row:
                        if not cell.tag.endswith("}c") and cell.tag != "c":
                            continue
                        cell_type = cell.attrib.get("t", "")
                        value_node = next((child for child in cell if child.tag.endswith("}v") or child.tag == "v"), None)
                        inline_node = next((child for child in cell if child.tag.endswith("}is") or child.tag == "is"), None)
                        value = ""
                        if cell_type == "s" and value_node is not None:
                            index = int(value_node.text or "-1")
                            value = shared[index] if 0 <= index < len(shared) else ""
                        elif cell_type == "inlineStr" and inline_node is not None:
                            value = "".join(node.text or "" for node in inline_node.iter() if node.tag.endswith("}t") or node.tag == "t")
                        elif value_node is not None:
                            value = value_node.text or ""
                        values.append(value)
                    if any(values):
                        rows.append("\t".join(values))
                        evidence.append({"id": f"sheet-{Path(sheet).stem}-row-{len(rows)}", "kind": "cell_row", "location": sheet, "text": "\t".join(values), "confidence": 1.0})
            return _normalize("\n".join(rows)), evidence
    except (OSError, KeyError, ValueError, ElementTree.ParseError, zipfile.BadZipFile) as exc:
        raise DocumentParseError("Excel 文档解析失败，请检查文件内容。") from exc


def _parse_swagger(content: bytes) -> tuple[str, list[dict[str, Any]]]:
    try:
        value = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise DocumentParseError("Swagger/OpenAPI 内容不是有效 JSON。") from exc
    if not isinstance(value, dict) or not (value.get("openapi") or value.get("swagger")):
        raise DocumentParseError("JSON 文档未声明 openapi 或 swagger 版本。")
    text = json.dumps(value, ensure_ascii=False, indent=2)
    evidence = []
    for path, methods in (value.get("paths") or {}).items():
        if isinstance(methods, dict):
            for method, operation in methods.items():
                if method.casefold() not in {"get", "post", "put", "patch", "delete", "options", "head", "trace"}:
                    continue
                summary = operation.get("summary") if isinstance(operation, dict) else ""
                evidence.append({"id": f"path-{len(evidence)+1}", "kind": "api_operation", "location": f"{method.upper()} {path}", "text": str(summary or path), "confidence": 1.0})
    return text, evidence


def _parse_image(content: bytes) -> tuple[str, list[dict[str, Any]], float, list[str]]:
    """Extract OCR text with confidence and fail closed when Chinese language data is absent."""
    try:
        from PIL import Image
        import pytesseract
        command = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
        if not command:
            candidate = Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "Tesseract-OCR" / "tesseract.exe"
            if candidate.is_file():
                command = str(candidate)
        if command:
            pytesseract.pytesseract.tesseract_cmd = command
        with Image.open(io.BytesIO(content)) as image:
            image = image.convert("RGB")
            image = image.resize((max(image.width, 1600), max(image.height, 900)))
            available = set(pytesseract.get_languages(config=""))
            requested = os.environ.get("TESSERACT_LANG", "chi_sim+eng")
            languages = "+".join(item for item in requested.split("+") if item in available)
            warnings: list[str] = []
            if "chi_sim" in requested.split("+") and "chi_sim" not in available:
                warnings.append("未安装中文 OCR 语言包（chi_sim），已停止输出低可信中文识别结果。")
                return "", [], 0.0, warnings
            if not languages:
                raise DocumentParseError("OCR 语言包不可用，请安装并配置 Tesseract 语言数据。")
            data = pytesseract.image_to_data(image, lang=languages, output_type=pytesseract.Output.DICT)
            words: list[str] = []
            evidence: list[dict[str, Any]] = []
            confidences: list[float] = []
            for index, raw in enumerate(data.get("text", [])):
                word = _normalize(str(raw))
                try:
                    confidence = float(data.get("conf", ["-1"])[index]) / 100.0
                except (TypeError, ValueError, IndexError):
                    confidence = 0.0
                if not word or confidence <= 0:
                    continue
                words.append(word)
                confidences.append(confidence)
                evidence.append({"id": f"ocr-{len(evidence)+1}", "kind": "ocr_word", "location": {"left": data.get("left", [0])[index], "top": data.get("top", [0])[index], "width": data.get("width", [0])[index], "height": data.get("height", [0])[index]}, "text": word, "confidence": round(confidence, 4)})
            text = _normalize(" ".join(words))
            confidence = round(sum(confidences) / len(confidences), 4) if confidences else 0.0
            return text, evidence, confidence, warnings
    except ImportError as exc:
        raise DocumentParseError("图片 OCR 依赖未安装，暂时无法解析图片。") from exc
    except Exception as exc:  # noqa: BLE001 - sanitize PIL/Tesseract failures
        raise DocumentParseError("图片 OCR 解析失败，请检查图片和 OCR 引擎。") from exc


def parse_document_bytes(content: bytes | bytearray, filename: str, *, source_url: str = "") -> ParsedDocument:
    """Parse an uploaded document by extension within the size and type boundary."""
    data = _bounded(bytes(content))
    suffix = Path(filename or "").suffix.casefold()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError("仅支持 PDF、Word、Excel、Markdown、Swagger、文本和图片文件。")
    try:
        if suffix in {".txt", ".md", ".markdown"}:
            text = _normalize(data.decode("utf-8-sig")); fmt = "markdown" if suffix != ".txt" else "text"
            evidence = [{"id": f"line-{index}", "kind": "line", "location": {"line": index}, "text": line, "confidence": 1.0} for index, line in enumerate(text.splitlines(), start=1) if line.strip()]
            confidence = 1.0; warnings = []
        elif suffix == ".pdf":
            pages = [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]
            text = _normalize("\n".join(pages)); fmt = "pdf"
            evidence = [{"id": f"page-{index}", "kind": "page", "location": {"page": index}, "text": _normalize(page), "confidence": 1.0} for index, page in enumerate(pages, start=1) if _normalize(page)]
            confidence = 1.0; warnings = []
        elif suffix == ".docx":
            document = WordDocument(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
            text = _normalize("\n".join(paragraphs + tables)); fmt = "word"
            evidence = [{"id": f"paragraph-{index}", "kind": "paragraph", "location": {"paragraph": index}, "text": _normalize(value), "confidence": 1.0} for index, value in enumerate(paragraphs + tables, start=1) if _normalize(value)]
            confidence = 1.0; warnings = []
        elif suffix == ".xlsx":
            text, evidence = _parse_xlsx(data); fmt = "excel"; confidence = 1.0; warnings = []
        elif suffix == ".json":
            text, evidence = _parse_swagger(data); fmt = "swagger"; confidence = 1.0; warnings = []
        else:
            text, evidence, confidence, warnings = _parse_image(data)
            fmt = "ocr"
    except DocumentParseError:
        raise
    except (OSError, UnicodeDecodeError, ValueError, TypeError) as exc:
        raise DocumentParseError("文档解析失败，请检查文件格式和内容。") from exc
    if not text:
        if warnings:
            raise DocumentParseError("；".join(warnings))
        raise DocumentParseError("文档未提取到可用文本。")
    return ParsedDocument(Path(filename).stem or "需求文档", text, fmt, "file", source_url, tuple(evidence), confidence, tuple(warnings))


def parse_file(path: str | Path) -> ParsedDocument:
    """Read and parse a local file after enforcing the size boundary."""
    file_path = Path(path)
    try:
        if not file_path.is_file() or file_path.stat().st_size > MAX_DOCUMENT_BYTES:
            raise DocumentParseError("文档不存在或超过10MB限制。")
        return parse_document_bytes(file_path.read_bytes(), file_path.name)
    except DocumentParseError:
        raise
    except OSError as exc:
        raise DocumentParseError("文档文件无法读取。") from exc


def parse_online(url: str) -> ParsedDocument:
    """Fetch through the existing provider adapter and normalize online content."""
    try:
        adapter = adapter_for(url)
        result: DocumentContent = adapter.fetch(url)
    except Exception as exc:  # noqa: BLE001 - hide network/provider details
        raise DocumentParseError("在线文档获取失败，请检查链接或稍后重试。") from exc
    fmt = "swagger" if isinstance(adapter, SwaggerAdapter) else "online"
    title = result.title or "在线文档"
    content = _normalize(result.content)
    if not content:
        raise DocumentParseError("在线文档未提取到可用文本。")
    evidence = [{"id": f"line-{index}", "kind": "online_line", "location": {"line": index}, "text": line, "confidence": 1.0} for index, line in enumerate(content.splitlines(), start=1) if line.strip()]
    return ParsedDocument(title, content, fmt, result.source_type, result.source_url, tuple(evidence), 1.0, ())


def parse_requirement_document(document: Any) -> ParsedDocument:
    """Parse a persisted requirement source and advance it to the analysis stage."""
    document.status = document.Status.PARSING
    document.save(update_fields=("status",))
    try:
        if document.source_type == document.SourceType.ONLINE_LINK:
            parsed = parse_online(document.source_url)
        elif document.source_type in {document.SourceType.FILE, document.SourceType.SCREENSHOT}:
            if document.file_path:
                parsed = parse_file(document.file_path)
            elif document.content_text.strip():
                parsed = parse_document_bytes(document.content_text.encode("utf-8"), f"{document.title}.md")
            else:
                raise DocumentParseError("文件来源必须提供文件或正文。")
        elif document.source_type == document.SourceType.MANUAL:
            parsed = parse_document_bytes(document.content_text.encode("utf-8"), f"{document.title}.md")
        else:
            raise DocumentParseError("暂不支持该需求文档来源类型。")
        document.content_text = parsed.content
        document.parse_evidence = list(parsed.evidence)
        document.parse_confidence = parsed.confidence
        document.parse_warnings = list(parsed.warnings)
        document.status = document.Status.ANALYZING
        document.save(update_fields=("content_text", "parse_evidence", "parse_confidence", "parse_warnings", "status"))
        return parsed
    except DocumentParseError as exc:
        document.status = document.Status.FAILED
        document.parse_warnings = [str(exc) or "需求来源解析失败，未生成可供分析的证据。"]
        document.save(update_fields=("status", "parse_warnings"))
        raise


__all__ = ["DocumentParseError", "MAX_DOCUMENT_BYTES", "ParsedDocument", "parse_document_bytes", "parse_file", "parse_online", "parse_requirement_document"]
